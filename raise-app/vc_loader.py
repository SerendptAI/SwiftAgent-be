"""
VC Contact Loader — reads the VC.xlsx workbook and exposes contacts by sheet.
"""

from pathlib import Path
from typing import List, Dict, Optional
import openpyxl

from config import config

# Column mapping for the VC spreadsheet
COLUMN_MAP = {
    0: "First Name",
    1: "Last Name",
    2: "Number",
    3: "Company",
    4: "City",
    5: "State",
    6: "Country",
    7: "Email",
    8: "Phone",
    9: "Website",
}


class VCLoader:
    """Load and parse the VC.xlsx workbook, with in-memory caching."""

    _cached_data = None
    _cached_summary = None

    def __init__(self, path: Optional[str] = None):
        self._path = Path(path or config.VC_FILE_PATH)
        if not self._path.is_absolute():
            self._path = Path(__file__).resolve().parent / self._path

    def _load_data(self):
        """Parse workbook into memory if not already cached."""
        if VCLoader._cached_data is not None:
            return

        import time
        t0 = time.time()
        wb = openpyxl.load_workbook(str(self._path), data_only=True, read_only=True)
        data = {}
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            contacts = []
            first_row = True

            for row in ws.iter_rows(values_only=True):
                if first_row:
                    first_row = False
                    continue  # skip header row

                contact = {}
                for idx, value in enumerate(row):
                    if idx < len(COLUMN_MAP):
                        key = COLUMN_MAP[idx]
                        contact[key] = str(value).strip() if value else ""

                # Skip completely empty rows
                if any(contact.get(k) for k in ["First Name", "Last Name", "Company", "Email"]):
                    first = contact.get("First Name", "").strip()
                    last = contact.get("Last Name", "").strip()
                    contact["Full Name"] = f"{first} {last}".strip()
                    contacts.append(contact)
            
            data[sheet_name] = contacts
        wb.close()
        
        VCLoader._cached_data = data

        # Precompute summary
        summary = {"sheets": [], "total_contacts": 0, "total_with_email": 0}
        for sheet, contacts in data.items():
            total = len(contacts)
            with_email = sum(1 for c in contacts if c.get("Email"))
            summary["sheets"].append({"name": sheet, "total": total, "with_email": with_email})
            summary["total_contacts"] += total
            summary["total_with_email"] += with_email
            
        VCLoader._cached_summary = summary
        print(f"Loaded VC.xlsx into memory in {time.time() - t0:.2f}s")

    def get_sheet_names(self) -> List[str]:
        """Return all sheet names from the workbook."""
        self._load_data()
        return list(VCLoader._cached_data.keys())

    def get_contacts(self, sheet_name: str) -> List[Dict]:
        """Return all contacts from a given sheet as list of dicts."""
        self._load_data()
        return VCLoader._cached_data.get(sheet_name, [])

    def get_all_contacts(self) -> Dict[str, List[Dict]]:
        """Return all contacts grouped by sheet."""
        self._load_data()
        return VCLoader._cached_data

    def get_summary(self) -> Dict:
        """Return a summary of all sheets and contact counts."""
        self._load_data()
        return VCLoader._cached_summary
